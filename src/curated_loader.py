"""
Curated Portfolio Loader and Validator for NCRB 2024 BDA Project.

Provides automated manifest loading, dataset verification, and clean DataFrame
extraction for both Tier 1 (District Core) and Tier 2 (Metropolitan Analytical)
tables.
"""

import json
import logging
from pathlib import Path
from typing import Dict, List, Optional, Tuple

import openpyxl
import pandas as pd

logging.basicConfig(level=logging.INFO, format="%(levelname)s: %(message)s")
logger = logging.getLogger(__name__)

PROJECT_ROOT = Path(__file__).resolve().parent.parent
MANIFEST_PATH = PROJECT_ROOT / "data" / "curated_portfolio_manifest.json"


def load_manifest() -> Dict:
    """Load the project dataset catalog manifest."""
    if not MANIFEST_PATH.exists():
        raise FileNotFoundError(f"Manifest not found at {MANIFEST_PATH}")
    with open(MANIFEST_PATH, "r", encoding="utf-8") as f:
        return json.load(f)


def verify_portfolio() -> pd.DataFrame:
    """Verify all 17 curated files on disk, inspect dimensions and health.

    Returns a summary DataFrame.
    """
    manifest = load_manifest()
    records = []

    for entry in manifest.get("datasets", []):
        rel_path = entry["relative_path"]
        abs_path = PROJECT_ROOT / rel_path
        exists = abs_path.exists()
        size_kb = round(abs_path.stat().st_size / 1024, 1) if exists else 0

        num_rows, num_cols, sheet_name = None, None, None
        status = "OK"

        if exists:
            try:
                wb = openpyxl.load_workbook(abs_path, read_only=True)
                sheet_name = wb.sheetnames[0]
                ws = wb[sheet_name]
                num_rows = ws.max_row
                num_cols = ws.max_column
                wb.close()
            except Exception as e:
                status = f"Read Error: {e}"
        else:
            status = "Missing File"

        records.append({
            "File ID": entry["file_id"],
            "Tier": entry["tier"],
            "Table Code": entry["table_code"],
            "File Name": entry["file_name"],
            "Exists": exists,
            "Size (KB)": size_kb,
            "Sheet": sheet_name,
            "Raw Rows": num_rows,
            "Raw Cols": num_cols,
            "Status": status,
        })

    return pd.DataFrame(records)


def clean_district_table(file_path: Path) -> pd.DataFrame:
    """Parse an NCRB district-level Excel file into a clean DataFrame.

    Strips multi-row hierarchical headers, identifies state boundaries, extracts
    district observations, and discards aggregate total rows.
    """
    wb = openpyxl.load_workbook(file_path, data_only=True)
    ws = wb[wb.sheetnames[0]]
    all_rows = list(ws.iter_rows(values_only=True))
    wb.close()

    if not all_rows:
        return pd.DataFrame()

    # Determine header rows (typically rows 1, 2, 3 in NCRB format)
    # Row 0 is typically the report title.
    # We construct column headers by combining non-empty values across rows 1, 2, 3
    col_names = []
    num_cols = len(all_rows[1]) if len(all_rows) > 1 else len(all_rows[0])

    for c in range(num_cols):
        parts = []
        for r in range(1, 4):
            if r < len(all_rows) and c < len(all_rows[r]):
                val = all_rows[r][c]
                if val is not None:
                    val_str = str(val).strip().replace("\n", " ")
                    if val_str and val_str not in parts and not val_str.isdigit():
                        parts.append(val_str)
        col_title = " - ".join(parts) if parts else f"Col_{c}"
        col_names.append(col_title)

    # Clean standardized column 0 and 1
    col_names[0] = "Serial_No"
    col_names[1] = "State_or_District"

    # Iterate over data rows (starting at row 5 or 6)
    clean_records = []
    current_state = "UNKNOWN"

    for r_idx in range(4, len(all_rows)):
        row = all_rows[r_idx]
        if not row or all(x is None for x in row):
            continue

        c0 = str(row[0]).strip() if row[0] is not None else ""
        c1 = str(row[1]).strip() if len(row) > 1 and row[1] is not None else ""

        # Check for State declaration header, e.g., 'State: Andhra Pradesh' or 'UT: Delhi'
        full_text = f"{c0} {c1}".strip()
        if "State:" in full_text or "UT:" in full_text or (c0.startswith("State") and not c1):
            state_raw = full_text.replace("State:", "").replace("UT:", "").strip()
            if state_raw:
                current_state = state_raw
            continue

        # Skip summary and total rows
        lowered = f"{c0} {c1}".lower()
        if any(term in lowered for term in ["total", "source:", "all-india", "all india"]):
            continue

        # Valid district row typically has an integer/float serial number or non-empty district name
        district_name = c1 if c1 else c0
        if not district_name or district_name.isdigit():
            continue

        row_dict = {"State": current_state, "District": district_name}
        for col_idx in range(2, min(len(col_names), len(row))):
            val = row[col_idx]
            # Convert numeric strings or numbers
            if val is None or str(val).strip() in ["-", "", "None", "NA", "N.A."]:
                row_dict[col_names[col_idx]] = 0
            else:
                try:
                    row_dict[col_names[col_idx]] = float(val) if "." in str(val) else int(val)
                except (ValueError, TypeError):
                    row_dict[col_names[col_idx]] = 0

        clean_records.append(row_dict)

    df = pd.DataFrame(clean_records)
    return df


def clean_metro_table(file_path: Path) -> pd.DataFrame:
    """Parse an NCRB metropolitan city Excel file into a clean DataFrame."""
    wb = openpyxl.load_workbook(file_path, data_only=True)
    ws = wb[wb.sheetnames[0]]
    all_rows = list(ws.iter_rows(values_only=True))
    wb.close()

    if not all_rows:
        return pd.DataFrame()

    # Find the header row (typically row 2 or 3)
    header_idx = 1
    for idx in range(min(5, len(all_rows))):
        row_str = " ".join([str(x).lower() for x in all_rows[idx] if x is not None])
        if "city" in row_str or "sl" in row_str:
            header_idx = idx
            break

    # Build column names
    col_names = []
    header_row = all_rows[header_idx]
    for c_idx, val in enumerate(header_row):
        name = str(val).strip().replace("\n", " ") if val is not None else f"Col_{c_idx}"
        col_names.append(name)

    records = []
    for r_idx in range(header_idx + 1, len(all_rows)):
        row = all_rows[r_idx]
        if not row or all(x is None for x in row):
            continue

        c0 = str(row[0]).strip() if row[0] is not None else ""
        c1 = str(row[1]).strip() if len(row) > 1 and row[1] is not None else ""

        # Skip total rows
        lowered = f"{c0} {c1}".lower()
        if "total" in lowered or "source" in lowered or "all-india" in lowered:
            continue

        city_name = c1 if c1 and not c1.isdigit() else c0
        if not city_name or city_name.isdigit():
            continue

        row_dict = {"City": city_name}
        for c_idx in range(2, min(len(col_names), len(row))):
            val = row[c_idx]
            col = col_names[c_idx]
            if val is None or str(val).strip() in ["-", "", "None", "NA", "N.A."]:
                row_dict[col] = 0
            else:
                try:
                    row_dict[col] = float(val) if "." in str(val) else int(val)
                except (ValueError, TypeError):
                    row_dict[col] = val

        records.append(row_dict)

    return pd.DataFrame(records)


if __name__ == "__main__":
    print("=" * 80)
    print("NCRB 2024 CURATED PORTFOLIO VERIFICATION REPORT")
    print("=" * 80)

    summary_df = verify_portfolio()
    print(summary_df[["File ID", "Tier", "File Name", "Size (KB)", "Raw Rows", "Raw Cols", "Status"]].to_string(index=False))

    print("\n" + "=" * 80)
    print("SAMPLE PARSING VALIDATION")
    print("=" * 80)

    # Test district parsing
    dist_sample = PROJECT_ROOT / "data" / "curated" / "district" / "1DistrictwiseIPCCrimes2024.xlsx"
    if dist_sample.exists():
        df_dist = clean_district_table(dist_sample)
        print(f"Parsed 1DistrictwiseIPCCrimes2024 -> Shape: {df_dist.shape}")
        print(f"Districts parsed: {len(df_dist)}, States found: {df_dist['State'].nunique()}")
        print(f"Sample states: {list(df_dist['State'].unique()[:5])}")

    # Test metro parsing
    metro_sample = PROJECT_ROOT / "data" / "curated" / "metropolitan" / "TABLE1B16.xlsx"
    if metro_sample.exists():
        df_metro = clean_metro_table(metro_sample)
        print(f"\nParsed TABLE1B16 (1B.1) -> Shape: {df_metro.shape}")
        print(f"Cities parsed: {len(df_metro)}")
        print(f"Sample cities: {list(df_metro['City'].unique()[:5])}")
