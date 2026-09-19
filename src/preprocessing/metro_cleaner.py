"""
Metropolitan Data Cleaner for NCRB 2024 Tier 2 Datasets.

Extracts city-level observations, parses multi-year columns (2022-2024),
rates, chargesheeting efficiency, and qualitative motives (murder & cybercrime).
"""

import re
import unicodedata
from pathlib import Path
from typing import Dict, List, Optional, Tuple

import openpyxl
import pandas as pd


def clean_city_name(name: str) -> str:
    """Normalize metropolitan city names."""
    if not name:
        return ""
    name = str(name).strip()
    name = re.sub(r"\s+", " ", name)
    name = re.sub(r"^\d+[\.\)]\s*", "", name)
    # Remove footnote markers like '+' or '*'
    name = re.sub(r"[\+\*]+$", "", name).strip()
    return name.title()


def clean_metro_generic_table(
    file_path: Path, prefix: str = "metro"
) -> pd.DataFrame:
    """Parse city-level tables with City in column index 1."""
    wb = openpyxl.load_workbook(file_path, data_only=True)
    ws = wb[wb.sheetnames[0]]
    all_rows = list(ws.iter_rows(values_only=True))
    wb.close()

    if not all_rows:
        return pd.DataFrame()

    # Find the header row (typically row 2)
    header_idx = 2
    for r_idx in range(min(5, len(all_rows))):
        row_str = " ".join([str(x).lower() for x in all_rows[r_idx] if x is not None])
        if "city" in row_str:
            header_idx = r_idx
            break

    # Build clean column headers by combining header rows
    num_cols = len(all_rows[header_idx])
    col_names = ["serial_no", "city"]

    for c in range(2, num_cols):
        parts = []
        for r in range(header_idx, min(header_idx + 2, len(all_rows))):
            val = all_rows[r][c] if c < len(all_rows[r]) else None
            if val is not None:
                v = str(val).strip().replace("\n", " ")
                # Only exclude bracketed index numbers like [1], [2], not 4-digit years like 2022
                if v and not re.match(r"^\[\d+\]$", v) and v not in parts:
                    parts.append(v)
        h_str = " ".join(parts) if parts else f"col_{c}"
        # Normalize
        h_clean = unicodedata.normalize("NFKD", h_str)
        h_clean = re.sub(r"[^\w\s-]", "", h_clean)
        h_clean = re.sub(r"[\s-]+", "_", h_clean).strip("_").lower()
        col_names.append(f"{prefix}_{h_clean}")

    clean_records = []
    for r_idx in range(header_idx + 1, len(all_rows)):
        row = all_rows[r_idx]
        if not row or all(x is None for x in row):
            continue

        c0 = str(row[0]).strip() if row[0] is not None else ""
        c1 = str(row[1]).strip() if len(row) > 1 and row[1] is not None else ""

        # Check for summary row
        lowered = f"{c0} {c1}".lower()
        if any(term in lowered for term in ["total", "source", "all-india", "table"]):
            continue

        city_raw = c1 if c1 and not re.match(r"^\[?\d+\]?$", c1) else c0
        if not city_raw or re.match(r"^\[?\d+\]?$", city_raw):
            continue

        norm_city = clean_city_name(city_raw)
        if not norm_city:
            continue

        record = {"city": norm_city}
        for c_idx in range(2, min(num_cols, len(row))):
            val = row[c_idx]
            col_key = col_names[c_idx]
            if val is None or str(val).strip() in ["-", "", "None", "NA", "N.A.", "*"]:
                record[col_key] = 0.0
            else:
                try:
                    record[col_key] = float(val) if "." in str(val) else int(val)
                except (ValueError, TypeError):
                    record[col_key] = 0.0

        clean_records.append(record)

    df = pd.DataFrame(clean_records)
    return df.drop_duplicates(subset=["city"]).reset_index(drop=True)


def parse_metro_juvenile_crime_demography(file_path: Path) -> pd.DataFrame:
    """Parse Table 5B.4 which is indexed by Crime Head rather than City."""
    wb = openpyxl.load_workbook(file_path, data_only=True)
    ws = wb[wb.sheetnames[0]]
    all_rows = list(ws.iter_rows(values_only=True))
    wb.close()

    records = []
    for r_idx in range(4, len(all_rows)):
        row = all_rows[r_idx]
        if not row or all(x is None for x in row):
            continue
        c0 = str(row[0]).strip() if row[0] is not None else ""
        c1 = str(row[1]).strip() if len(row) > 1 and row[1] is not None else ""
        if "total" in f"{c0} {c1}".lower() or "source" in f"{c0} {c1}".lower():
            continue
        crime_head = c1 if c1 else c0
        if not crime_head or re.match(r"^\[?\d+\]?$", crime_head):
            continue

        rec = {"crime_head": crime_head.strip()}
        # Extract numeric columns safely
        for c_idx, val in enumerate(row[2:], start=2):
            col_name = f"col_{c_idx}"
            if val is None or str(val).strip() in ["-", "", "NA"]:
                rec[col_name] = 0
            else:
                try:
                    rec[col_name] = int(val) if isinstance(val, (int, float)) else int(float(str(val).strip()))
                except (ValueError, TypeError):
                    rec[col_name] = 0
        records.append(rec)

    return pd.DataFrame(records)


def build_metro_master_trends(curated_dir: Path) -> pd.DataFrame:
    """Build a unified 3-year metropolitan comparative dataset across IPC, Women, and Children."""
    df_ipc = clean_metro_generic_table(curated_dir / "TABLE1B16.xlsx", prefix="ipc")
    df_burden = clean_metro_generic_table(curated_dir / "TABLE1B33.xlsx", prefix="total")
    df_women = clean_metro_generic_table(curated_dir / "TABLE3B13.xlsx", prefix="women")
    df_child = clean_metro_generic_table(curated_dir / "TABLE4B13.xlsx", prefix="child")
    df_murder = clean_metro_generic_table(curated_dir / "TABLE2B24.xlsx", prefix="murder")
    df_cyber = clean_metro_generic_table(curated_dir / "TABLE9B33.xlsx", prefix="cyber")
    df_socio = clean_metro_generic_table(curated_dir / "TABLE5B63.xlsx", prefix="juv_socio")

    # Merge on city
    master_metro = df_ipc.copy()
    for df in [df_burden, df_women, df_child, df_murder, df_cyber, df_socio]:
        master_metro = pd.merge(master_metro, df, on="city", how="left")

    master_metro = master_metro.fillna(0)
    return master_metro
